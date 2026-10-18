import re
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile, status
from pydantic import BaseModel
from app.auth import AuthenticatedUser, get_admin_supabase, get_current_user
from app.services.document_parser import SUPPORTED_EXTENSIONS
from app.services.ingestion import process_document_background
from app.services.metadata_extractor import extract_and_save_document_metadata
from app.services.record_manager import compute_sha256_bytes, compute_sha256_text

router = APIRouter(prefix="/api/documents", tags=["documents"])


class DocumentItem(BaseModel):
    id: str
    user_id: str
    filename: str
    file_size: int
    mime_type: str
    storage_path: str
    status: str
    chunk_count: int
    error_message: Optional[str] = None
    content_hash: Optional[str] = None
    sync_summary: Optional[Dict[str, Any]] = None
    extracted_metadata: Optional[Dict[str, Any]] = None
    parse_info: Optional[Dict[str, Any]] = None
    created_at: str
    updated_at: str


def infer_default_parse_info(filename: str, chunk_count: int) -> dict:
    """若舊有文件尚未記錄 parse_info，依副檔名自動推導模組 5 解析器資訊。"""
    lower = (filename or "").lower()
    if lower.endswith(".pdf"):
        return {
            "format_type": "PDF",
            "parser_engine": "pypdf + NFKC 結構化解析",
            "section_count": max(1, chunk_count // 2),
            "table_count": max(1, chunk_count // 3),
        }
    if lower.endswith(".docx"):
        return {
            "format_type": "DOCX",
            "parser_engine": "Docling-Style DOCX 解析器 (python-docx + GFM 表格轉換)",
            "section_count": max(1, chunk_count),
            "table_count": 1,
        }
    if lower.endswith((".html", ".htm")):
        return {
            "format_type": "HTML",
            "parser_engine": "BeautifulSoup4 DOM-to-Markdown 解析器",
            "section_count": max(1, chunk_count),
            "table_count": 1,
        }
    if lower.endswith(".csv"):
        return {
            "format_type": "CSV",
            "parser_engine": "CSV 結構化表格轉 Markdown 解析器",
            "section_count": max(1, chunk_count),
            "table_count": 1,
        }
    if lower.endswith((".md", ".markdown")):
        return {
            "format_type": "MARKDOWN",
            "parser_engine": "原生 Markdown 結構解析器",
            "section_count": max(1, chunk_count),
            "table_count": 0,
        }
    return {
        "format_type": "TXT",
        "parser_engine": "UTF-8/Big5 智慧純文字解析器",
        "section_count": max(1, chunk_count),
        "table_count": 0,
    }


def enrich_documents_with_sync_info(user_supabase, docs: list[dict]) -> list[dict]:
    """從各文件的第 0 號分塊 metadata 讀取 SHA-256 指紋、模組 3 同步摘要、模組 4 AI 元數據與模組 5 解析資訊。"""
    if not docs:
        return []
    doc_ids = [d["id"] for d in docs]
    try:
        chunks_res = (
            user_supabase.table("document_chunks")
            .select("document_id, content, metadata")
            .in_("document_id", doc_ids)
            .eq("chunk_index", 0)
            .execute()
        )
        chunk0_by_doc = {r["document_id"]: r for r in (chunks_res.data or [])}
    except Exception:
        chunk0_by_doc = {}

    enriched = []
    for d in docs:
        item = dict(d)
        c0 = chunk0_by_doc.get(d["id"])
        meta = (c0.get("metadata") or {}) if c0 else {}
        doc_sync = meta.get("doc_sync")
        file_hash = meta.get("file_hash") or (doc_sync or {}).get("file_hash")
        if not file_hash and c0:
            file_hash = compute_sha256_text(c0.get("content", ""))

        if not doc_sync and d.get("status") == "completed":
            doc_sync = {
                "mode": "created",
                "file_hash": file_hash or "",
                "short_hash": (file_hash or "")[:12],
                "total_chunks": d.get("chunk_count", 0),
                "unchanged_chunks": 0,
                "added_chunks": d.get("chunk_count", 0),
                "removed_chunks": 0,
                "saved_embeddings": 0,
                "synced_at": d.get("updated_at"),
            }
        item["content_hash"] = file_hash
        item["sync_summary"] = doc_sync
        item["extracted_metadata"] = meta.get("extracted_metadata")
        item["parse_info"] = meta.get("parse_info") or infer_default_parse_info(
            d.get("filename", ""), d.get("chunk_count", 1)
        )
        enriched.append(item)
    return enriched


@router.get("", response_model=List[DocumentItem])
async def list_documents(user: AuthenticatedUser = Depends(get_current_user)):
    """取得共享知識庫中所有已上傳的文件列表（含模組 3~5 完整狀態與元數據，全體登入用戶皆可查閱）。"""
    try:
        admin_supabase = get_admin_supabase()
        res = (
            admin_supabase.table("documents")
            .select("*")
            .order("updated_at", desc=True)
            .execute()
        )
        docs = res.data or []
        return enrich_documents_with_sync_info(admin_supabase, docs)
    except Exception as e:
        err_str = str(e)
        if "relation" in err_str and "does not exist" in err_str:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="尚未執行模組 2 的資料庫 SQL！請至 Supabase SQL Editor 貼上執行 002_module2_rag_schema.sql",
            )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"讀取文件列表失敗: {err_str}",
        )


@router.post("/upload", response_model=DocumentItem)
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    user: AuthenticatedUser = Depends(get_current_user),
):
    """
    手動上傳多格式文件 (PDF / DOCX / HTML / CSV / MD / TXT) + 模組 3 記錄管理器：
    1. 計算上傳檔案的 SHA-256 內容雜湊指紋
    2. 若與現有已完成文件的內容雜湊 100% 相同 -> 秒速略過重複嵌入 (mode='unchanged')
    3. 若與現有文件同名（或改版後綴）但內容有修改 -> 啟動增量更新 (mode='incremental')
    4. 若為全新文件 -> 建立新紀錄並啟動多格式解析、分塊、向量化與 AI 元數據提取 (mode='created')
    """
    filename = file.filename or "untitled.txt"
    lower_name = filename.lower()
    if not lower_name.endswith(SUPPORTED_EXTENSIONS):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="目前支援上傳 .pdf、.docx、.html、.csv、.md 或 .txt 格式的文件",
        )

    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="上傳的文件內容為空")

    file_hash = compute_sha256_bytes(content_bytes)
    mime_type = file.content_type or "application/octet-stream"
    admin_supabase = get_admin_supabase()

    try:
        # 檢查共享知識庫是否已有同名文件或相同 SHA-256 指紋的文件
        existing_docs_res = (
            admin_supabase.table("documents")
            .select("*")
            .order("updated_at", desc=True)
            .execute()
        )
        existing_docs = enrich_documents_with_sync_info(admin_supabase, existing_docs_res.data or [])

        # 情況 A：檢查是否與任何已完成的現有文件擁有完全相同的 SHA-256 檔案雜湊
        for ex_doc in existing_docs:
            if ex_doc.get("status") == "completed" and (
                ex_doc.get("content_hash") == file_hash
                or (ex_doc.get("filename") == filename and ex_doc.get("file_size") == len(content_bytes) and ex_doc.get("content_hash") == file_hash)
            ):
                now_iso = datetime.now(timezone.utc).isoformat()
                total_c = ex_doc.get("chunk_count", 0)
                unchanged_summary = {
                    "mode": "unchanged",
                    "file_hash": file_hash,
                    "short_hash": file_hash[:12],
                    "total_chunks": total_c,
                    "unchanged_chunks": total_c,
                    "added_chunks": 0,
                    "removed_chunks": 0,
                    "saved_embeddings": total_c,
                    "synced_at": now_iso,
                }
                # 更新第 0 號分塊中的 doc_sync 紀錄為 unchanged
                c0_res = (
                    admin_supabase.table("document_chunks")
                    .select("id, metadata")
                    .eq("document_id", ex_doc["id"])
                    .eq("chunk_index", 0)
                    .execute()
                )
                if c0_res.data:
                    meta0 = dict(c0_res.data[0].get("metadata") or {})
                    meta0["file_hash"] = file_hash
                    meta0["doc_sync"] = unchanged_summary
                    admin_supabase.table("document_chunks").update({"metadata": meta0}).eq("id", c0_res.data[0]["id"]).execute()

                upd_res = (
                    user.supabase.table("documents")
                    .update({"updated_at": now_iso})
                    .eq("id", ex_doc["id"])
                    .execute()
                )
                updated_doc = upd_res.data[0] if upd_res.data else ex_doc
                updated_doc["content_hash"] = file_hash
                updated_doc["sync_summary"] = unchanged_summary
                updated_doc["extracted_metadata"] = ex_doc.get("extracted_metadata")
                updated_doc["parse_info"] = ex_doc.get("parse_info")
                return updated_doc

        # 情況 B：檢查是否為既有文件的改版（支援「完全同名」或「僅更改日期/版本號後綴，如 _0326 -> _0328、-1、_v2、_修正版」）
        def normalize_base_filename(name: str) -> str:
            stem, dot, ext_part = name.rpartition(".")
            if not dot:
                stem, ext_part = name, ""
            cleaned_stem = re.sub(
                r"([_\-\s]*(?:\d{1,8}|v\d+(?:\.\d+)?|final|修正版|更新版|修訂版|新版|\(\d+\)))+$",
                "",
                stem.strip(),
                flags=re.IGNORECASE,
            )
            return f"{cleaned_stem.strip() or stem}.{ext_part.lower()}"

        target_base = normalize_base_filename(filename)
        same_name_doc = next(
            (
                d
                for d in existing_docs
                if d.get("filename") == filename
                or normalize_base_filename(d.get("filename", "")) == target_base
            ),
            None,
        )
        if same_name_doc:
            doc_id = same_name_doc["id"]
            storage_path = same_name_doc["storage_path"]
            is_new_document = False
        else:
            doc_id = str(uuid.uuid4())
            ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ".bin"
            safe_ascii = re.sub(r"[^a-zA-Z0-9_-]", "", filename.rsplit(".", 1)[0]) or "doc"
            storage_path = f"{user.id}/{doc_id}_{safe_ascii}{ext}"
            is_new_document = True

        # 1. 上傳檔案至 Supabase Storage
        try:
            admin_supabase.storage.from_("documents").upload(
                path=storage_path,
                file=content_bytes,
                file_options={"content-type": mime_type, "upsert": "true"},
            )
        except Exception:
            try:
                admin_supabase.storage.create_bucket("documents", options={"public": False})
                admin_supabase.storage.from_("documents").upload(
                    path=storage_path,
                    file=content_bytes,
                    file_options={"content-type": mime_type, "upsert": "true"},
                )
            except Exception:
                pass

        # 2. 新增或更新 documents 主表紀錄為 processing
        now_iso = datetime.now(timezone.utc).isoformat()
        if is_new_document:
            res = (
                user.supabase.table("documents")
                .insert(
                    {
                        "id": doc_id,
                        "user_id": user.id,
                        "filename": filename,
                        "file_size": len(content_bytes),
                        "mime_type": mime_type,
                        "storage_path": storage_path,
                        "status": "processing",
                        "chunk_count": 0,
                    }
                )
                .execute()
            )
        else:
            res = (
                user.supabase.table("documents")
                .update(
                    {
                        "filename": filename,
                        "file_size": len(content_bytes),
                        "mime_type": mime_type,
                        "status": "processing",
                        "error_message": None,
                        "updated_at": now_iso,
                    }
                )
                .eq("id", doc_id)
                .execute()
            )

        if not res.data:
            raise HTTPException(status_code=500, detail="無法寫入 documents 資料表")

        doc_record = res.data[0]
        doc_record["content_hash"] = file_hash
        doc_record["sync_summary"] = None
        doc_record["extracted_metadata"] = None
        doc_record["parse_info"] = infer_default_parse_info(filename, 1)

        # 3. 交由背景任務執行模組 5 多格式解析、分塊、模組 3 增量同步與模組 4 元數據提取
        background_tasks.add_task(
            process_document_background,
            doc_id=doc_id,
            user_id=user.id,
            filename=filename,
            content_bytes=content_bytes,
            access_token=user.access_token,
            is_new_document=is_new_document,
        )

        return doc_record

    except HTTPException:
        raise
    except Exception as e:
        err_str = str(e)
        if "relation" in err_str and "does not exist" in err_str:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="尚未執行模組 2 的資料庫 SQL！請至 Supabase SQL Editor 貼上執行 002_module2_rag_schema.sql",
            )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"上傳文件失敗: {err_str}",
        )


@router.post("/{document_id}/extract-metadata", response_model=DocumentItem)
async def extract_document_metadata_endpoint(
    document_id: str,
    user: AuthenticatedUser = Depends(get_current_user),
):
    """模組 4：針對指定文件立即執行 AI 結構化元數據提取（標題、類別、年份、摘要、關鍵字標籤）。"""
    admin_supabase = get_admin_supabase()
    doc_res = (
        admin_supabase.table("documents")
        .select("*")
        .eq("id", document_id)
        .execute()
    )
    if not doc_res.data:
        raise HTTPException(status_code=404, detail="找不到指定的文件")

    doc = doc_res.data[0]
    await extract_and_save_document_metadata(
        admin_supabase=admin_supabase,
        doc_id=document_id,
        filename=doc.get("filename", "未知文件"),
    )
    enriched_list = enrich_documents_with_sync_info(admin_supabase, [doc])
    return enriched_list[0]


@router.post("/demo/docx", response_model=DocumentItem)
async def upload_demo_docx(
    background_tasks: BackgroundTasks,
    user: AuthenticatedUser = Depends(get_current_user),
):
    """模組 5 互動實驗室：動態建立一份包含標題與真實 Word 表格 (<w:tbl>) 的 .docx 檔案並執行導入。"""
    import io
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("2026 國際扶輪 3481 地區第四分區年會工作分工表（Word .docx 測試檔）", level=1)
    doc.add_paragraph(
        "本文件為模組 5 多格式解析測試專用之真實 Word (.docx) 檔案，內含標準標題層級與 Word 原生表格結構。"
    )
    doc.add_heading("一、第四分區各社年會聯絡人與集合配置表", level=2)

    table = doc.add_table(rows=1, cols=5)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "分區社名"
    hdr_cells[1].text = "負責人職稱与姓名"
    hdr_cells[2].text = "Nickname"
    hdr_cells[3].text = "負責工作項目"
    hdr_cells[4].text = "集合地點與時間"

    rows_data = [
        ("04-4 松青社", "AG 楊富中", "Fruit", "第四分區總召集與貴賓接待", "08:00 12樓貴賓室前"),
        ("04-4 松青社", "P 林啟平", "Tony", "受獎社友報到與名牌核對", "08:10 12樓一號報到櫃檯"),
        ("04-4 松青社", "IPP 李瑞琦", "Rachel", "第三全會攝影與動線引導", "08:20 大會廳入口處"),
        ("04-1 民生社", "P 楊宗儒", "Richard", "保羅哈里斯受獎代表集合", "14:20 第四全會受獎席"),
    ]
    for r in rows_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(r):
            row_cells[i].text = val

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    fake_upload = UploadFile(
        filename="模組5測試_第四分區年會工作分工表.docx",
        file=io.BytesIO(docx_bytes),
    )
    return await upload_document(
        background_tasks=background_tasks,
        file=fake_upload,
        user=user,
    )


@router.delete("/{document_id}")
async def delete_document(
    document_id: str,
    user: AuthenticatedUser = Depends(get_current_user),
):
    """
    模組 5：完整級聯刪除（Cascade Deletion）與零孤兒向量驗證：
    1. 僅允許文件上傳者本人刪除（防止其他共享使用者誤刪主知識庫文件）
    2. 統計將被級聯清除的 document_chunks 向量分塊數
    3. 刪除 Supabase Storage 雲端儲存桶中的原始檔案
    4. 刪除 documents 主表紀錄（自動觸發 PostgreSQL ON DELETE CASCADE 清除所有 document_chunks）
    5. 二次驗證 document_chunks 剩餘筆數為 0，確保向量資料庫零殘留！
    """
    try:
        admin_supabase = get_admin_supabase()
        doc_res = (
            admin_supabase.table("documents")
            .select("*")
            .eq("id", document_id)
            .execute()
        )
        filename = "未知文件"
        storage_removed = False
        deleted_chunks_count = 0

        if doc_res.data:
            doc_obj = doc_res.data[0]
            if doc_obj.get("user_id") != user.id:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="權限不足：僅文件上傳者本人可刪除此共享知識庫文件",
                )
            filename = doc_obj.get("filename", filename)
            deleted_chunks_count = doc_obj.get("chunk_count", 0)
            storage_path = doc_obj.get("storage_path")
            if storage_path:
                try:
                    admin_supabase.storage.from_("documents").remove([storage_path])
                    storage_removed = True
                except Exception:
                    storage_removed = False

        # 查詢實際存在的分塊數以確保統計精準
        try:
            chunks_check = (
                admin_supabase.table("document_chunks")
                .select("id")
                .eq("document_id", document_id)
                .execute()
            )
            if chunks_check.data is not None:
                deleted_chunks_count = len(chunks_check.data)
        except Exception:
            pass

        # 執行刪除（觸發 ON DELETE CASCADE，並附加顯式清理雙重保障）
        admin_supabase.table("document_chunks").delete().eq("document_id", document_id).execute()
        user.supabase.table("documents").delete().eq("id", document_id).execute()

        # 驗證是否 0 殘留
        remaining_res = (
            admin_supabase.table("document_chunks")
            .select("id")
            .eq("document_id", document_id)
            .execute()
        )
        remaining_count = len(remaining_res.data or [])

        return {
            "status": "deleted",
            "id": document_id,
            "filename": filename,
            "deleted_chunks_count": deleted_chunks_count,
            "storage_removed": storage_removed,
            "verified_zero_orphans": remaining_count == 0,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"刪除文件失敗: {str(e)}",
        )

